from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"
OUT = DOCS / "前端团队AI学习分享_2026年9月.docx"
MAP = ASSETS / "knowledge-map.png"
FONT = "/System/Library/Fonts/PingFang.ttc"

BLUE = "1473E6"
INK = "15171A"
MUTED = "626A73"
LINE = "D9DEE5"
PALE_BLUE = "EEF5FF"
PALE_GREEN = "EAF8F1"
PALE_AMBER = "FFF4DF"
WHITE = "FFFFFF"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=LINE, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_run_font(run, size=None, bold=None, color=None, name="PingFang SC"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    run_pr.append(c)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "PingFang SC")
    run_pr.append(fonts)
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(doc, text, level=1, number=None):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    if number:
        run = p.add_run(f"{number}  {text}")
    else:
        run = p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc, title, body, color=PALE_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, 11, True, INK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run(body)
    set_run_font(r2, 10, False, MUTED)


def add_source_line(doc, index, title, url, note=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(f"[{index}] ")
    set_run_font(r, 8.5, True, MUTED)
    add_hyperlink(p, title, url)
    if note:
        r2 = p.add_run(f" — {note}")
        set_run_font(r2, 8.5, False, MUTED)


def create_map():
    ASSETS.mkdir(parents=True, exist_ok=True)
    scale = 160
    width, height = int(13.2 * scale), int(6.8 * scale)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(FONT, 27)
    subtitle_font = ImageFont.truetype(FONT, 15)
    node_title_font = ImageFont.truetype(FONT, 18)
    node_detail_font = ImageFont.truetype(FONT, 14)
    nodes = {
        "base": (0.4, 4.8, 2.2, 1.25, "20/80 基础", "ML 统计直觉\nLLM 原理"),
        "coding": (0.4, 2.7, 2.2, 1.25, "协作开发", "规格 · 上下文\n审查 · 测试"),
        "python": (0.4, 0.6, 2.2, 1.25, "Python 工程", "异步 API\n测试 · 超时"),
        "context": (3.1, 4.05, 2.4, 1.25, "上下文工程", "模型 API · Schema\n选择 · 压缩"),
        "rag": (6.0, 4.85, 2.35, 1.25, "知识供给", "RAG · 数据 · 权限\n检索 · 引用"),
        "agent": (6.0, 2.95, 2.35, 1.25, "Agent 系统", "Workflow · Tool\nState · MCP"),
        "ui": (6.0, 1.05, 2.35, 1.25, "AI UI / UX", "流式 · 多模态\n证据 · 用户控制"),
        "eval": (8.85, 4.85, 2.25, 1.25, "评测", "黄金集 · Grader\nTrace · 归因"),
        "reliability": (8.85, 2.95, 2.25, 1.25, "可靠交付", "队列 · 幂等 · 回滚\nLLMOps · 成本"),
        "security": (8.85, 1.05, 2.25, 1.25, "安全", "最小权限 · 确认\n审计 · 沙箱"),
        "product": (11.55, 2.75, 1.35, 1.65, "结果", "可用、可测\n可控的\nAI 产品"),
    }

    colors = {
        "base": "#F4F5F7", "coding": "#EEF5FF", "python": "#EEF5FF",
        "context": "#EEF5FF", "rag": "#F0F6FF", "agent": "#F0F6FF",
        "ui": "#FFF4DF", "eval": "#EAF8F1", "reliability": "#EAF8F1",
        "security": "#FFF0F0", "product": "#1473E6",
    }
    edges = [
        ("base", "context"), ("coding", "context"), ("python", "rag"),
        ("context", "rag"), ("context", "agent"), ("context", "ui"),
        ("rag", "agent"), ("rag", "eval"), ("agent", "eval"),
        ("agent", "reliability"), ("agent", "security"), ("ui", "product"),
        ("eval", "product"), ("reliability", "product"), ("security", "product"),
    ]
    for source, target in edges:
        sx, sy, sw, sh, *_ = nodes[source]
        tx, ty, tw, th, *_ = nodes[target]
        start = (int((sx + sw) * scale), int((6.8 - (sy + sh / 2)) * scale))
        end = (int(tx * scale), int((6.8 - (ty + th / 2)) * scale))
        draw.line([start, end], fill="#AAB2BD", width=2)
        ex, ey = end
        draw.polygon([(ex, ey), (ex - 11, ey - 6), (ex - 11, ey + 6)], fill="#AAB2BD")
    for key, (x, y, w, h, title, detail) in nodes.items():
        x1 = int(x * scale)
        y1 = int((6.8 - y - h) * scale)
        x2 = int((x + w) * scale)
        y2 = int((6.8 - y) * scale)
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=colors[key], outline="#D6DCE3", width=2)
        text_color = "white" if key == "product" else "#15171A"
        draw.text((x1 + 24, y1 + 20), title, font=node_title_font, fill=text_color)
        draw.multiline_text((x1 + 24, y1 + 60), detail, font=node_detail_font,
                            fill=("#EAF2FF" if key == "product" else "#606873"), spacing=7)
    draw.text((64, 40), "前端 → AI 长期能力图谱", font=title_font, fill="#15171A")
    draw.text((660, 49), "框架会换，连接关系和失败边界更值得长期学习", font=subtitle_font, fill="#6A727C")
    image.save(MAP, quality=95)


def build_document():
    create_map()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.64)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title_style = styles["Title"]
    title_style.font.name = "PingFang SC"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    title_style.font.color.rgb = RGBColor.from_string(INK)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, color, before, after in (
        ("Heading 1", 21, INK, 8, 10),
        ("Heading 2", 14.5, INK, 11, 6),
        ("Heading 3", 11.5, INK, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "PingFang SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Cover
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(78)
    r = p.add_run("AI 迭代这么快\n前端还值得学什么")
    set_run_font(r, 30, True, INK)
    p.paragraph_format.line_spacing = 1.05

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("从当前模型能力 北京真实 JD 和 50 组面经反推学习路线")
    set_run_font(r, 14, False, INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(195)
    r = p.add_run("前端团队内部分享")
    set_run_font(r, 12, True, INK)
    p2 = doc.add_paragraph()
    r2 = p2.add_run("调研校准日期  2026-09-06")
    set_run_font(r2, 10, False, MUTED)

    add_callout(
        doc,
        "先给结论",
        "不要赌某个模型、框架或提示词技巧能活多久。长期投资“定义问题—组织上下文—连接数据和工具—评测结果—可靠交付—让人可控”这条完整链路。",
        PALE_BLUE,
    )
    doc.add_page_break()

    # Executive summary
    add_heading(doc, "分享提纲", 1)
    summary = doc.add_table(rows=3, cols=2)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.autofit = False
    for row in summary.rows:
        row.cells[0].width = Inches(1.35)
        row.cells[1].width = Inches(5.45)
        for cell in row.cells:
            set_cell_margins(cell, 150, 160, 150, 160)
    items = [
        ("01  背景", "AI 能写更多代码，但“能生成”不等于“能交付”；解释调研口径与证据边界。"),
        ("02  知识图谱", "哪些是核心投入、哪些保留 20/80、哪些按岗位选学、哪些降低优先级。"),
        ("03  开源项目", "展示如何用 JD 和面经反推学习内容，并把路线变成面试与练习闭环。"),
    ]
    for index, (label, detail) in enumerate(items):
        set_cell_shading(summary.cell(index, 0), PALE_BLUE if index == 0 else "F6F7F9")
        r1 = summary.cell(index, 0).paragraphs[0].add_run(label)
        set_run_font(r1, 10, True, BLUE if index == 0 else INK)
        r2 = summary.cell(index, 1).paragraphs[0].add_run(detail)
        set_run_font(r2, 10, False, MUTED)

    add_heading(doc, "核心判断", 2)
    add_bullets(doc, [
        "核心投入：AI 协作开发、Python、上下文工程、RAG 数据链路、Agent / Workflow、评测、AI UI / UX、后端可靠性、安全与 LLMOps。",
        "保留 20/80：机器学习与统计直觉、LLM / Transformer 原理。",
        "按岗位选学：微调与开源模型、Next.js 等特定全栈框架、高级多模态 / GPU 管线。",
        "降低优先级：提示词魔法句式、背框架 API、朴素 RAG Demo、多 Agent 炫技、每日追模型榜单。",
    ])

    add_heading(doc, "证据不是一个数字", 2)
    add_body(doc, "精确 JD 记录职位详情、职位 ID、发布日期和核验日期；招聘雷达只是公司官网核验任务；面经频次表示有多少来源覆盖该主题，并不声称某句话被逐字问了多少次。")
    add_callout(doc, "边界说明", "当前精确 JD 抽样以能稳定访问的百度北京职位为主；50 组面经包含个人自述和题单汇总。它们能提示方向，但不能代表完整劳动力市场。", PALE_AMBER)
    doc.add_page_break()

    # 1 Background
    add_heading(doc, "背景", 1, "01")
    add_heading(doc, "一个月后再看：变化很快，但路线没有整体失效", 2)
    add_body(doc, "模型完成软件任务的跨度继续变长，简单页面和局部实现更容易被自动化；与此同时，真实开发者仍普遍遇到“几乎正确”的输出、调试负担和信任问题。前端工作没有消失，它正在从“亲手写每一行”转向“把问题说清、让模型看懂仓库、验证结果、守住体验和风险”。")

    evidence_table = doc.add_table(rows=1, cols=3)
    evidence_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    evidence_table.autofit = False
    headers = ["实际信号", "说明", "对学习的含义"]
    widths = [1.55, 2.65, 2.65]
    for i, text in enumerate(headers):
        cell = evidence_table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, INK)
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(text)
        set_run_font(r, 9, True, WHITE)
    set_repeat_table_header(evidence_table.rows[0])
    rows = [
        ("代码 Agent", "Anthropic 编程交互样本中自动化占比高，JS / HTML 与 UI / UX 常见。", "简单实现会贬值，复杂交互、性能和验证升值。"),
        ("开发者调查", "Stack Overflow 2025：不信任高于信任，“几乎正确”与调试负担普遍。", "必须学评测、审查和失败复现。"),
        ("长任务", "METR 与 OpenAI 都观察到模型承担更长任务的趋势，但真实任务更脏。", "规格、状态、恢复和人工确认成为基本功。"),
        ("语言生态", "GitHub 2025：TypeScript 位居首位，Python 第二。", "保留类型化前端优势，同时补 Python。"),
    ]
    for row_index, row_data in enumerate(rows):
        cells = evidence_table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if row_index % 2:
                set_cell_shading(cells[i], "F8F9FA")
            r = cells[i].paragraphs[0].add_run(value)
            set_run_font(r, 8.7, i == 0, INK if i == 0 else MUTED)

    add_heading(doc, "不要只听预测：用 JD 和面经做复核", 2)
    add_body(doc, "厂商对未来的共识是开发者会转向理解、指导、验证，人的判断、品味和责任更重要；但厂商有推广产品的利益，预测只能当方向假设。我们把它们与当前 JD 和面经交叉验证后再决定优先级。")
    add_body(doc, "本轮北京样本里，Python、模型 API、RAG、Agent、工具、评测、系统设计、稳定性和安全反复出现。面经则把应用岗、平台岗和算法岗的学习深度拉开：应用岗看 RAG 质量、工具容错、上下文和项目价值；平台岗看状态、协议与系统设计；算法岗才显著增加训练深度。")

    doc.add_page_break()
    add_heading(doc, "反推方法", 2)
    method = doc.add_table(rows=4, cols=2)
    method.alignment = WD_TABLE_ALIGNMENT.CENTER
    method.autofit = False
    steps = [
        ("1  收 JD", "优先官方详情，保留职位 ID、日期、城市、状态和原链接。"),
        ("2  抽能力", "把 RAG / Agent 拆成数据、评测、工具、状态、权限、成本等能力。"),
        ("3  看面经", "统计来源覆盖主题，不把一个汇总帖拆成很多“独立样本”。"),
        ("4  变练习", "能力映射到题目目的、参考答案、闪卡、脑图与专项任务。"),
    ]
    for i, (label, value) in enumerate(steps):
        method.rows[i].cells[0].width = Inches(1.2)
        method.rows[i].cells[1].width = Inches(5.6)
        set_cell_shading(method.rows[i].cells[0], PALE_BLUE)
        for cell in method.rows[i].cells:
            set_cell_margins(cell, 120, 150, 120, 150)
        r1 = method.rows[i].cells[0].paragraphs[0].add_run(label)
        set_run_font(r1, 9.3, True, BLUE)
        r2 = method.rows[i].cells[1].paragraphs[0].add_run(value)
        set_run_font(r2, 9.3, False, MUTED)
    # 2 Knowledge graph
    add_heading(doc, "值得学习的内容知识图谱", 1, "02")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(MAP), width=Inches(6.85))
    cap = doc.add_paragraph("图 1  长期能力的依赖关系：框架会换，连接关系与失败边界更稳定")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.runs[0], 8.5, False, MUTED)

    add_heading(doc, "核心投入", 2)
    core_rows = [
        ("AI 协作开发", "规格、仓库上下文、任务拆解、diff 审查、测试、回滚", "可执行任务包 + 审查记录"),
        ("Python 工程", "类型、pytest、asyncio、FastAPI、超时与重试", "异步模型网关"),
        ("上下文工程", "模型 API、Schema、上下文选择 / 压缩、缓存、路由", "固定评测集 + 版本记录"),
        ("RAG 知识供给", "解析、混合检索、重排、权限、引用、增量索引", "能分层归因的 RAG"),
        ("Agent 系统", "Workflow、工具契约、状态、检查点、人工确认、MCP", "可恢复、可审计的 Agent"),
        ("评测与 Trace", "黄金集、grader、错误分类、质量 / 延迟 / 成本门槛", "自动回归报告"),
        ("AI UI / UX", "流式、多模态、长任务状态、证据、撤销、用户控制", "五状态交互原型"),
        ("可靠性 / 安全 / LLMOps", "队列、幂等、回滚、最小权限、审计、灰度、成本", "发布清单 + 威胁模型"),
    ]
    core = doc.add_table(rows=1, cols=3)
    core.alignment = WD_TABLE_ALIGNMENT.CENTER
    core.autofit = False
    widths = [1.45, 3.55, 1.8]
    for i, label in enumerate(("知识块", "够用标准", "可验证产出")):
        core.rows[0].cells[i].width = Inches(widths[i])
        set_cell_shading(core.rows[0].cells[i], INK)
        set_cell_margins(core.rows[0].cells[i])
        r = core.rows[0].cells[i].paragraphs[0].add_run(label)
        set_run_font(r, 8.8, True, WHITE)
    set_repeat_table_header(core.rows[0])
    for index, row in enumerate(core_rows):
        cells = core.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_margins(cells[i], 85, 100, 85, 100)
            if index % 2:
                set_cell_shading(cells[i], "F8F9FA")
            r = cells[i].paragraphs[0].add_run(value)
            set_run_font(r, 8.2, i == 0, INK if i == 0 else MUTED)

    add_heading(doc, "保留 20/80 与按岗位选学", 2)
    decisions = doc.add_table(rows=1, cols=4)
    decisions.alignment = WD_TABLE_ALIGNMENT.CENTER
    decisions.autofit = False
    headers = ["判断", "知识块", "学习边界", "为什么"]
    widths = [1.0, 1.45, 2.55, 1.8]
    for i, text in enumerate(headers):
        decisions.rows[0].cells[i].width = Inches(widths[i])
        set_cell_shading(decisions.rows[0].cells[i], INK)
        set_cell_margins(decisions.rows[0].cells[i])
        r = decisions.rows[0].cells[i].paragraphs[0].add_run(text)
        set_run_font(r, 8.8, True, WHITE)
    set_repeat_table_header(decisions.rows[0])
    rows = [
        ("继续学习", "ML / 统计直觉", "训练 / 验证 / 测试、过拟合、数据泄漏、核心指标、实验与误差分析", "应用岗要会判断数据与指标，不必先推完整公式"),
        ("继续学习", "LLM 原理", "token、attention 直觉、上下文、采样、KV Cache、训练 / 对齐区别", "能连接到质量、延迟与成本即可"),
        ("按岗位选", "微调 / 开源模型", "先做 Prompt、RAG、基座升级对照；算法 / 平台 / 私有化岗再深学", "数据与评测不成熟时，微调容易浪费"),
        ("按岗位选", "Next.js", "理解服务端边界、缓存和密钥安全；框架深度跟目标岗位走", "它是实现选项，不是 AI 门票"),
        ("按岗位选", "多模态 / GPU", "AI 体验、多模态应用、创意工具岗深学", "其他岗位先理解输入输出、时延与成本"),
    ]
    for index, row in enumerate(rows):
        cells = decisions.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_margins(cells[i], 100, 105, 100, 105)
            if index < 2:
                set_cell_shading(cells[i], PALE_BLUE if i == 0 else "F9FBFF")
            else:
                set_cell_shading(cells[i], PALE_AMBER if i == 0 else "FFFCF6")
            r = cells[i].paragraphs[0].add_run(value)
            set_run_font(r, 8.2, i <= 1, INK if i <= 1 else MUTED)

    add_heading(doc, "降低优先级", 2)
    low_rows = [
        ("提示词魔法句式", "上下文工程 + 结构化输出 + 固定评测"),
        ("背框架 API", "状态、工具契约、检查点、恢复等原理"),
        ("朴素 RAG Demo", "数据治理 + 混合检索 + 引用 + 分层评测"),
        ("多 Agent 炫技", "先用确定性 Workflow 或单 Agent"),
        ("把 Next.js 当前置课", "按 JD 选栈，先补 Python / 评测 / 可靠性"),
        ("每日追模型榜单", "每月用真实任务更新一次能力基线"),
    ]
    for old, replacement in low_rows:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{old}：")
        set_run_font(r1, 10, True, INK)
        r2 = p.add_run(replacement)
        set_run_font(r2, 10, False, MUTED)

    add_heading(doc, "12 周务实路线", 2)
    route = [
        ("1—2 周", "AI 协作开发", "规格、计划、审查、测试、回滚"),
        ("3—4 周", "Python 服务", "FastAPI、asyncio、超时、重试、日志"),
        ("5—6 周", "上下文工程", "结构化输出 + 20 条固定评测"),
        ("7—8 周", "可信 RAG", "引用、权限、检索 / 生成分层评测"),
        ("9—10 周", "Agent 可靠性", "工具、状态、幂等、恢复、人工确认"),
        ("11—12 周", "AI 体验与发布", "流式 UI、Trace、安全回归、项目复盘"),
    ]
    route_table = doc.add_table(rows=1, cols=3)
    route_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    route_table.autofit = False
    for i, value in enumerate(("时间", "主题", "交付重点")):
        set_cell_shading(route_table.rows[0].cells[i], BLUE)
        set_cell_margins(route_table.rows[0].cells[i])
        r = route_table.rows[0].cells[i].paragraphs[0].add_run(value)
        set_run_font(r, 9, True, WHITE)
    for week, topic, output in route:
        cells = route_table.add_row().cells
        for i, value in enumerate((week, topic, output)):
            set_cell_margins(cells[i], 95, 120, 95, 120)
            r = cells[i].paragraphs[0].add_run(value)
            set_run_font(r, 8.8, i < 2, INK if i < 2 else MUTED)

    add_callout(doc, "执行原则", "每两周只交一个可演示、可测量的小产出。不要同时开五门课，也不要把看新闻算作完成学习。", PALE_GREEN)
    # 3 Open source project
    add_heading(doc, "网站开源项目", 1, "03")
    add_heading(doc, "它不是课程目录，而是一条学习闭环", 2)
    add_body(doc, "JD 告诉我们企业愿意为什么能力付钱；面经告诉我们企业怎样验证；行业研究帮助判断知识会不会快速折旧；训练工具再把“知道”变成“能回答、能做、能复盘”。")

    features = doc.add_table(rows=4, cols=2)
    features.alignment = WD_TABLE_ALIGNMENT.CENTER
    features.autofit = False
    feature_rows = [
        ("学习路线", "按重要度、难度、天数和本轮判断排序；知识块含入门 / 进阶 / 精通边界与 6 条资源。"),
        ("岗位机会", "精确 JD 与招聘检索分层；保留职位 ID、抓取日期和原链接，不虚构职位充数量。"),
        ("面试训练", "50 个来源聚合高频问题，展示考察目的、两种参考答案、语音 / 文本回答、评分与进步曲线。"),
        ("刻意练习", "按知识块提供低中高闪卡、脑图、键盘操作和能提交复盘的专项任务。"),
    ]
    for i, (name, detail) in enumerate(feature_rows):
        features.rows[i].cells[0].width = Inches(1.35)
        features.rows[i].cells[1].width = Inches(5.45)
        set_cell_shading(features.rows[i].cells[0], PALE_BLUE)
        for cell in features.rows[i].cells:
            set_cell_margins(cell, 125, 150, 125, 150)
        r1 = features.rows[i].cells[0].paragraphs[0].add_run(name)
        set_run_font(r1, 9.5, True, BLUE)
        r2 = features.rows[i].cells[1].paragraphs[0].add_run(detail)
        set_run_font(r2, 9.5, False, MUTED)

    add_heading(doc, "这次开源前完成的整理", 2)
    add_bullets(doc, [
        "保留原有 Git 历史，没有重建仓库。",
        "增加 MIT License、贡献说明、环境变量示例和项目 README。",
        "把 2026-09 判断写回学习路线，并新增 AI 协作开发知识块、6 条资源、6 张闪卡和专项任务。",
        "更新北京精确 JD 样本和 2026 年 8 月面经来源。",
        "明确证据口径：招聘雷达不是独立 JD，薪资未公开就不猜，面经频次是来源覆盖度。",
        "DeepSeek API Key 只从服务端环境变量读取，仓库不包含真实 Key。",
    ])

    add_heading(doc, "现场演示路径", 2)
    demo = [
        "在学习路线筛选“核心投入”，打开 Agent / Workflow，展示判断依据与来源。",
        "在岗位机会打开精确 JD，展示职位 ID、发布日期和证据层级。",
        "在面试训练按热度选题，先看题目目的，再口述 60 秒并评分。",
        "打开回答历史看分数曲线，把薄弱项送到刻意练习。",
        "用脑图说明知识之间的依赖，而不是把它当一张课程清单。",
    ]
    for index, item in enumerate(demo, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.38)
        p.paragraph_format.first_line_indent = Inches(-0.38)
        r1 = p.add_run(f"{index:02d}  ")
        set_run_font(r1, 10, True, BLUE)
        r2 = p.add_run(item)
        set_run_font(r2, 10, False, INK)

    add_heading(doc, "适合团队共同维护", 2)
    add_bullets(doc, [
        "每月核验失效 JD、新岗位与抓取日期。",
        "给面经补来源、日期和岗位类型，避免重复样本。",
        "提交更好的中英文学习资源和真实项目练习题。",
        "用同一组题跑不同模型，维护质量、延迟和成本基线。",
        "改进无障碍、移动端、语音识别和数据导入导出。",
    ])

    add_callout(doc, "最后的判断", "“AI 迭代太快，学了也白学”只对一半。模型名称、框架 API 和花哨 Demo 会快速折旧；定义目标、组织上下文、连接真实数据、验证结果、控制风险和做好体验，反而会随着 AI 变强而更重要。", PALE_BLUE)
    doc.add_page_break()

    # Sources
    add_heading(doc, "调研来源", 1)
    add_body(doc, "以下来源均在 2026-09-06 前后核验。厂商研究存在自身立场，招聘与面经样本也有选择偏差，因此正文只做交叉判断，不把单一来源当确定预测。")
    sources = [
        ("Anthropic｜软件开发中的 AI 使用", "https://www.anthropic.com/research/impact-software-development", "50 万次交互，观察自动化与 UI / Web 使用"),
        ("Stack Overflow｜2025 AI 调查", "https://survey.stackoverflow.co/2025/ai", "信任、准确性与调试负担"),
        ("Stack Overflow｜调查方法", "https://survey.stackoverflow.co/2025/methodology", "49,009 份有效样本"),
        ("METR｜Task time horizons", "https://metr.org/time-horizons/", "软件任务跨度趋势及限制"),
        ("OpenAI｜How agents are transforming work", "https://openai.com/index/how-agents-are-transforming-work/", "长任务使用的抽样分析"),
        ("GitHub｜Octoverse 2025", "https://github.blog/news-insights/octoverse/octoverse-a-new-developer-joins-github-every-second-as-ai-leads-typescript-to-1/", "TypeScript、Python 与 AI 项目趋势"),
        ("GitHub｜开发者新身份", "https://github.blog/news-insights/octoverse/the-new-identity-of-a-developer-what-changes-and-what-doesnt-in-the-ai-era/", "理解、指导、验证"),
        ("Anthropic｜Building effective agents", "https://www.anthropic.com/engineering/building-effective-agents", "简单架构、工具与评测"),
        ("Anthropic｜Context engineering", "https://www.anthropic.com/engineering/effective-context-engineering-for-ai-agents", "从提示词到上下文系统"),
        ("Anthropic｜Agent evals", "https://www.anthropic.com/engineering/demystifying-evals-for-ai-agents", "task、trial、grader、trace 与 outcome"),
        ("OWASP｜Agentic Top 10 2026", "https://genai.owasp.org/resource/owasp-top-10-for-agentic-applications-for-2026/", "自主 Agent 风险"),
        ("MCP｜2025-11 授权规范", "https://modelcontextprotocol.io/specification/2025-11-25/basic/authorization", "协议授权边界"),
        ("Stanford｜AI Index 2026", "https://hai.stanford.edu/assets/files/ai_index_report_2026.pdf", "能力、基准与饱和风险"),
        ("百度 J100679", "https://talent.baidu.com/jobs/detail/GRADUATE/66a12645-f0f1-435c-8426-9fb91f1be330", "Python、RAG、Agent、API、Docker / K8s"),
        ("百度 J103341", "https://talent.baidu.com/jobs/detail/SOCIAL/a5ff8d15-b547-4a87-ba55-a128dae953cd", "API、CLI、Skill 与 Agent-friendly 基础设施"),
        ("百度 J101017", "https://talent.baidu.com/jobs/detail/GRADUATE/02f73086-be71-4d09-8d6e-f1c6981b8b48", "Agent 设计与评测体系"),
        ("百度 J102217", "https://talent.baidu.com/jobs/detail/SOCIAL/f16d38a1-440b-4e7b-b09b-ddfdfaf643e4", "训练、应用、评测、RAG 与多模态"),
        ("牛客｜腾讯 / 百度大模型面经", "https://www.nowcoder.com/discuss/878600528970735616", "Python、检索、工具、系统设计与项目深挖"),
        ("牛客｜2026 应用岗位复盘", "https://www.nowcoder.com/discuss/914178628659707904", "RAG、Agent、上下文、MCP、安全与工程化"),
        ("牛客｜Agent 岗位分层复盘", "https://www.nowcoder.com/discuss/916347378695692288", "应用、平台与算法岗位差异"),
    ]
    for index, (title, url, note) in enumerate(sources, 1):
        if index == 11:
            doc.add_page_break()
        add_source_line(doc, index, title, url, note)

    # Headers and footers
    for table in doc.tables:
        set_table_borders(table)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for sec in doc.sections:
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("Frontend → AI · 2026-09")
        set_run_font(hr, 8, False, INK)
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("前端团队内部分享  ·  证据可回溯，结论可更新")
        set_run_font(run, 8, False, MUTED)

    doc.core_properties.title = "AI 迭代这么快，前端还值得学什么"
    doc.core_properties.subject = "2026-09 前端转 AI 学习路线与开源项目分享"
    doc.core_properties.author = "Frontend AI Compass"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
